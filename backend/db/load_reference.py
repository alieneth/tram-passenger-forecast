"""DATA-1: загрузка справочника (depot, route, stop, route_stop) в PostgreSQL.

Источник — единственный xlsx организаторов (data/Хакатон_справочники_трамвай_10_маршрутов.xlsx),
поля берём как есть из листов, ничего не придумываем (см. правило в CLAUDE.md).

Маршруты проекта — CLAUDE.md раздел 4: 1, 5, 7, 11, 12, 17, 25, 26, 28, 50. Из них в справочнике
(лист "Маршруты GTFS_ROUTES") есть только 1, 5, 7, 11, 12 — у них реальные route_long_name, depot_id,
координаты остановок. Для 17, 25, 26, 28, 50 в справочнике нет ничего (has_geometry = false).
route 5 — is_new = true (в train/test нет ни одной строки, см. CLAUDE.md раздел 4 и 13.2).

Запуск: DATABASE_URL=postgresql://tram:change_me@localhost:5433/tram python backend/db/load_reference.py
"""

import os
from datetime import date, datetime

import openpyxl
import psycopg

PROJECT_ROUTES = [1, 5, 7, 11, 12, 17, 25, 26, 28, 50]
NO_HISTORY_ROUTES = {5}
XLSX_PATH = "data/Хакатон_справочники_трамвай_10_маршрутов.xlsx"
# Единственное депо во всём справочнике (лист "Наряд": depot_id 133 -> depot_name).
DEPOT_NAMES = {133: "Трамвайное управление"}


def parse_date(value: str | None) -> date | None:
    return datetime.strptime(value, "%Y-%m-%d").date() if value else None


def load_routes_sheet(wb: openpyxl.Workbook) -> dict[int, dict]:
    ws = wb["Маршруты GTFS_ROUTES"]
    by_short_name = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        (_route_id, _reg_num, route_short_name, route_long_name, *_rest, depot_id) = row
        by_short_name[int(route_short_name)] = {
            "route_long_name": route_long_name,
            "depot_id": int(depot_id) if depot_id else None,
            "date_start": parse_date(row[8]),
        }
    return by_short_name


def load_stops_sheet(wb: openpyxl.Workbook) -> list[dict]:
    ws = wb["Остановки GTFS_STOPS"]
    stops = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        (stop_id, stop_name, _desc, stop_lat, stop_lon, street, region, district, *_rest, pavilion, _actual, _deleted) = row
        stops.append(
            {
                "stop_id": int(stop_id),
                "stop_name": stop_name,
                "stop_lat": float(stop_lat),
                "stop_lon": float(stop_lon),
                "street": street,
                "district": district,
                "region": region,
                "has_pavilion": pavilion == "1",
            }
        )
    return stops


def load_route_stop_sheet(wb: openpyxl.Workbook, wanted_routes: set[int]) -> list[dict]:
    ws = wb["Порядок_с_координатами"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        # 0 route_id, 1 route_short_name, 4 trip_id, 6 direction_id, 9 stop_sequence, 10 stop_id, 12 stop_mode
        route_short_name = row[1]
        trip_id = row[4]
        direction_id = row[6]
        stop_sequence = row[9]
        stop_id = row[10]
        stop_mode = row[12]
        route = int(route_short_name)
        if route not in wanted_routes:
            continue
        rows.append(
            {
                "route": route,
                "stop_id": int(stop_id),
                "trip_id": int(trip_id),
                "direction_id": int(direction_id),
                "stop_sequence": int(stop_sequence),
                "stop_mode": int(stop_mode) if stop_mode not in (None, "") else None,
            }
        )
    return rows


def build_routes(spravochnik: dict[int, dict]) -> list[dict]:
    routes = []
    for route in PROJECT_ROUTES:
        info = spravochnik.get(route)
        routes.append(
            {
                "route": route,
                "route_long_name": info["route_long_name"] if info else None,
                "depot_id": info["depot_id"] if info else None,
                "date_start": info["date_start"] if info else None,
                "is_new": route in NO_HISTORY_ROUTES,
                "has_geometry": info is not None,
            }
        )
    return routes


def main() -> None:
    database_url = os.environ["DATABASE_URL"]
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)

    spravochnik = load_routes_sheet(wb)
    stops = load_stops_sheet(wb)
    routes = build_routes(spravochnik)
    # Геометрия есть только у маршрутов проекта, которые нашлись в справочнике (1, 5, 7, 11, 12).
    geometry_routes = {r["route"] for r in routes if r["has_geometry"]}
    route_stops = load_route_stop_sheet(wb, geometry_routes)

    depot_ids = {r["depot_id"] for r in routes if r["depot_id"]}

    with psycopg.connect(database_url) as conn, conn.cursor() as cur:
        for depot_id in depot_ids:
            cur.execute(
                "INSERT INTO depot (depot_id, depot_name) VALUES (%s, %s) "
                "ON CONFLICT (depot_id) DO NOTHING",
                (depot_id, DEPOT_NAMES[depot_id]),
            )
        for r in routes:
            cur.execute(
                "INSERT INTO route (route, route_long_name, depot_id, date_start, is_new, has_geometry) "
                "VALUES (%(route)s, %(route_long_name)s, %(depot_id)s, %(date_start)s, %(is_new)s, %(has_geometry)s) "
                "ON CONFLICT (route) DO UPDATE SET "
                "route_long_name = EXCLUDED.route_long_name, depot_id = EXCLUDED.depot_id, "
                "date_start = EXCLUDED.date_start, is_new = EXCLUDED.is_new, has_geometry = EXCLUDED.has_geometry",
                r,
            )
        for s in stops:
            cur.execute(
                "INSERT INTO stop (stop_id, stop_name, stop_lat, stop_lon, street, district, region, has_pavilion) "
                "VALUES (%(stop_id)s, %(stop_name)s, %(stop_lat)s, %(stop_lon)s, %(street)s, %(district)s, %(region)s, %(has_pavilion)s) "
                "ON CONFLICT (stop_id) DO UPDATE SET "
                "stop_name = EXCLUDED.stop_name, stop_lat = EXCLUDED.stop_lat, stop_lon = EXCLUDED.stop_lon, "
                "street = EXCLUDED.street, district = EXCLUDED.district, region = EXCLUDED.region, "
                "has_pavilion = EXCLUDED.has_pavilion",
                s,
            )
        cur.execute("DELETE FROM route_stop")
        cur.executemany(
            "INSERT INTO route_stop (route, stop_id, trip_id, direction_id, stop_sequence, stop_mode) "
            "VALUES (%(route)s, %(stop_id)s, %(trip_id)s, %(direction_id)s, %(stop_sequence)s, %(stop_mode)s)",
            route_stops,
        )
        conn.commit()

    print(
        f"depot={len(depot_ids)} route={len(routes)} stop={len(stops)} "
        f"route_stop={len(route_stops)} (геометрия: {sorted(geometry_routes)})"
    )


if __name__ == "__main__":
    main()
